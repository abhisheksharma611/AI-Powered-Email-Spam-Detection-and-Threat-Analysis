# -*- coding: utf-8 -*-
"""Export the SAME Phase-2 CONTENT to an editable Word file for teammates.

Output: ../Phase2_Report_Contents.docx
Run: python export_docx.py  (from sp2phase2/)
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

import build_phase2_report as B

OUT = os.path.join(os.path.dirname(HERE), "Phase2_Report_Contents.docx")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

ACCENT = RGBColor(0xC5, 0x5A, 0x11)

title = doc.add_heading("AI-Powered Email Spam Detection and Threat Analysis (2026-2027)", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("Phase-2 Report Contents  \u2014  SP2 scope (editable draft)")
run.bold = True
run.font.size = Pt(13)

note = doc.add_paragraph()
run = note.add_run("TEAM NOTE \u2014 delete this box before submission: ")
run.bold = True
run.font.color.rgb = ACCENT
note.add_run(
    "This file holds the full Phase-2 text: Chapters 1\u20136 (41 sections), "
    "13 tables, 3 figures (embedded below their sections), and References. "
    "Chapter 6 is a test PLAN (all rows say Planned) \u2014 do not mark anything Passed. "
    "No Results chapter, no screenshots, no Conclusion/Future Scope in Phase 2. "
    "Add your college title/certificate pages at the top before printing."
)

CHAPTER_TITLES = {}
for it in B.CONTENT:
    if it[0] == "chapter":
        CHAPTER_TITLES[it[1]] = it[2]

n_fig = n_tab = 0
for it in B.CONTENT:
    kind = it[0]
    if kind == "chapter":
        doc.add_heading("Chapter %s: %s" % (it[1], it[2]), level=1)
    elif kind == "sec":
        level = 3 if it[1].count(".") == 2 else 2
        doc.add_heading("%s  %s" % (it[1], it[2]), level=level)
    elif kind == "body":
        doc.add_paragraph(it[1])
    elif kind == "table":
        headers, rows = it[1], it[2]
        cap = it[3] if len(it) > 3 else ""
        t = doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        for j, h in enumerate(headers):
            cell = t.rows[0].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(h))
            run.bold = True
        for i, r in enumerate(rows, start=1):
            for j, c in enumerate(r):
                t.rows[i].cells[j].text = str(c)
        if cap:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(cap)
            run.italic = True
            run.font.size = Pt(10)
        n_tab += 1
    elif kind == "diag":
        doc.add_picture(it[1], width=Inches(6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(it[2])
        run.italic = True
        run.font.size = Pt(10)
        n_fig += 1
    elif kind == "references":
        doc.add_heading("References", level=1)
        for ref_list in it[1]:
            for k, ref in enumerate(ref_list, start=1):
                doc.add_paragraph("[%d] %s" % (k, ref))

doc.save(OUT)
print("DOCX written:", OUT, "| figures:", n_fig, "| tables:", n_tab)
